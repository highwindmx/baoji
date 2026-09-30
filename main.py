"""宝记 后端：FastAPI + SQLite(WAL)。

局域网宝宝数据收集系统：喂养/睡眠/尿布/生长曲线(WHO 百分位)/体温/疫苗/用药/里程碑/提醒/导出。
启动：python main.py  （支持 --host / --port 0 自动分配 / --no-auth）
"""
import os
import re
import json
import shutil
import socket
import hashlib
import uuid
import io
from datetime import datetime, date, timedelta, timezone

from fastapi import (
    FastAPI, Depends, HTTPException, UploadFile, File, APIRouter, Form, Header,
)
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, Response
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy import select, func, and_
from sqlalchemy.orm import Session

import database as db
from database import (
    Baby, FeedingSession, FeedingLog, Sleep, Diaper, Medication, Temperature,
    Vaccine, Growth, Milestone, Photo, get_db, UPLOAD_DIR, BASE_DIR,
    ensure_default_baby, get_single_baby_id,
)
import schemas as s

# ---------------- 配置（环境变量优先） ----------------
from dotenv import load_dotenv
load_dotenv()
PASSWORD = os.environ.get("BABYWEB_PASSWORD", "baby1234")
SECRET = os.environ.get("BABYWEB_SECRET", "babyweb-local-secret")
EXPECTED_TOKEN = hashlib.sha256((PASSWORD + SECRET).encode()).hexdigest()
NO_AUTH = os.environ.get("BABYWEB_NOAUTH", "0").lower() in ("1", "true", "yes")
DEFAULT_BABY_ID = 1  # 启动后由 ensure_default_baby 修正

APP_NAME = "宝记"

# 允许上传的媒体格式（照片墙/里程碑/尿布照片：图片 + 视频）
IMAGE_EXTS = (".jpg", ".jpeg", ".png", ".gif", ".webp", ".heic")
VIDEO_EXTS = (".mp4", ".mov", ".m4v", ".webm", ".avi", ".3gp")
MEDIA_EXTS = IMAGE_EXTS + VIDEO_EXTS

FEED_TYPE_LABEL = {"breast": "母乳", "formula": "奶粉", "mixed": "混合"}

# 中国免疫规划（0~24 月）参考接种计划，用于疫苗提醒。
# month = 应接种月龄（相对出生）；name 用于与用户接种记录匹配“已完成”。
VACCINE_SCHEDULE = [
    {"name": "卡介苗(BCG)", "month": 0},
    {"name": "乙肝第1针", "month": 0},
    {"name": "乙肝第2针", "month": 1},
    {"name": "脊灰第1剂", "month": 2},
    {"name": "脊灰第2剂", "month": 3},
    {"name": "百白破第1剂", "month": 3},
    {"name": "脊灰第3剂", "month": 4},
    {"name": "百白破第2剂", "month": 4},
    {"name": "百白破第3剂", "month": 5},
    {"name": "乙肝第3针", "month": 6},
    {"name": "A群流脑第1剂", "month": 6},
    {"name": "麻腮风第1剂", "month": 8},
    {"name": "乙脑减毒第1剂", "month": 8},
    {"name": "A群流脑第2剂", "month": 9},
    {"name": "百白破加强", "month": 18},
    {"name": "麻腮风加强", "month": 18},
    {"name": "甲肝疫苗", "month": 18},
    {"name": "乙脑加强", "month": 24},
]


def make_token() -> str:
    return EXPECTED_TOKEN


def require_token(x_token: str = Header(None)) -> str:
    # 免验证模式（--no-auth / BABYWEB_NOAUTH=1）下直接放行。
    if NO_AUTH:
        return x_token or "noauth"
    if x_token and x_token == EXPECTED_TOKEN:
        return x_token
    raise HTTPException(status_code=401, detail="未授权，请先登录")


# ---------------- App & 中间件 ----------------
app = FastAPI(title="宝记 宝宝数据收集", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

api = APIRouter(dependencies=[Depends(require_token)])


# ---------------- 登录 / 配置（公开） ----------------
@app.post("/api/login")
def login(password: str = Form(...)):
    if NO_AUTH:
        return {"token": "noauth"}
    if hashlib.sha256((password + SECRET).encode()).hexdigest() != EXPECTED_TOKEN:
        raise HTTPException(status_code=401, detail="口令错误")
    return {"token": make_token()}


@app.get("/api/config")
def config():
    return {"app_name": APP_NAME, "auth_required": not NO_AUTH, "version": "1.0.0"}


@app.get("/api/me")
def me(token: str = Depends(require_token)):
    return {"ok": True}


# ---------------- 单宝宝（GET/PUT） ----------------
@api.get("/baby", response_model=s.BabyOut)
def get_baby(d: Session = Depends(get_db)):
    b = d.get(Baby, DEFAULT_BABY_ID)
    if not b:
        b = ensure_default_baby()
    return b


@api.put("/baby", response_model=s.BabyOut)
def update_baby(body: s.BabyCreate, d: Session = Depends(get_db)):
    b = d.get(Baby, DEFAULT_BABY_ID) or ensure_default_baby()
    for k, v in body.model_dump().items():
        setattr(b, k, v)
    d.commit()
    d.refresh(b)
    return b


# ---------------- 通用 CRUD 工厂 ----------------
MODELS = {
    "sleep": (Sleep, s.SleepCreate, s.SleepOut),
    "diaper": (Diaper, s.DiaperCreate, s.DiaperOut),
    "medication": (Medication, s.MedicationCreate, s.MedicationOut),
    "temperature": (Temperature, s.TemperatureCreate, s.TemperatureOut),
    "vaccine": (Vaccine, s.VaccineCreate, s.VaccineOut),
    "growth": (Growth, s.GrowthCreate, s.GrowthOut),
    "milestone": (Milestone, s.MilestoneCreate, s.MilestoneOut),
}


def register_crud(prefix, model, create_schema, out_schema):
    tag = prefix.capitalize()

    @api.get(f"/{prefix}s", response_model=list[out_schema], tags=[tag])
    def list_items(
        baby_id: int = None,
        from_: str = None,
        to_: str = None,
        d: Session = Depends(get_db),
    ):
        if baby_id is None:
            baby_id = DEFAULT_BABY_ID
        q = select(model).where(model.baby_id == baby_id)
        time_col = None
        for cand in ("start_time", "time", "record_date", "date", "created_at"):
            if hasattr(model, cand):
                time_col = getattr(model, cand)
                break
        if time_col is not None and from_:
            q = q.where(time_col >= from_)
        if time_col is not None and to_:
            q = q.where(time_col <= to_)
        if time_col is not None:
            q = q.order_by(time_col.desc())
        else:
            q = q.order_by(model.id.desc())
        return d.scalars(q).all()

    @api.post(f"/{prefix}s", response_model=out_schema, tags=[tag])
    def create_item(body: create_schema, d: Session = Depends(get_db)):
        data = body.model_dump()
        if data.get("baby_id") is None:
            data["baby_id"] = DEFAULT_BABY_ID
        obj = model(**data)
        d.add(obj)
        d.commit()
        d.refresh(obj)
        # 尿布建档后把记录 id 回写进媒体 .json 说明文件
        if model is Diaper and obj.photo_filename:
            _write_meta(obj.photo_filename, kind="diaper", record_id=obj.id)
        return obj

    @api.put(f"/{prefix}s/{{item_id}}", response_model=out_schema, tags=[tag])
    def update_item(item_id: int, body: create_schema, d: Session = Depends(get_db)):
        obj = d.get(model, item_id)
        if not obj:
            raise HTTPException(404, "记录不存在")
        data = body.model_dump()
        data.pop("baby_id", None)   # 不允许改动记录归属
        for k, val in data.items():
            setattr(obj, k, val)
        d.commit()
        d.refresh(obj)
        return obj

    @api.delete(f"/{prefix}s/{{item_id}}", tags=[tag])
    def delete_item(item_id: int, d: Session = Depends(get_db)):
        obj = d.get(model, item_id)
        if not obj:
            raise HTTPException(404, "记录不存在")
        if model is Diaper and obj.photo_filename:
            _remove_media(obj.photo_filename)
        d.delete(obj)
        d.commit()
        return {"ok": True}


for _p, (_m, _c, _o) in MODELS.items():
    register_crud(_p, _m, _c, _o)


# ---------------- 喂养（一场 = 开始/结束/备注 + 多种奶量明细） ----------------
def _feeding_payload(sess_obj: "FeedingSession", d: Session) -> dict:
    logs = d.scalars(
        select(FeedingLog).where(FeedingLog.session_id == sess_obj.id).order_by(FeedingLog.id)
    ).all()
    return {
        "id": sess_obj.id, "baby_id": sess_obj.baby_id,
        "start_time": sess_obj.start_time, "end_time": sess_obj.end_time,
        "note": sess_obj.note or "", "created_at": sess_obj.created_at,
        "logs": [
            {"id": l.id, "feed_type": l.feed_type, "amount_ml": l.amount_ml}
            for l in logs
        ],
    }


@api.get("/feedings", response_model=list[s.FeedingSessionOut], tags=["Feeding"])
def list_feedings(
    baby_id: int = None, from_: str = None, to_: str = None,
    d: Session = Depends(get_db),
):
    baby_id = baby_id or DEFAULT_BABY_ID
    q = select(FeedingSession).where(FeedingSession.baby_id == baby_id)
    if from_:
        q = q.where(FeedingSession.start_time >= from_)
    if to_:
        q = q.where(FeedingSession.start_time <= to_)
    q = q.order_by(FeedingSession.start_time.desc())
    items = d.scalars(q).all()
    return [_feeding_payload(x, d) for x in items]


@api.post("/feedings", response_model=s.FeedingSessionOut, tags=["Feeding"])
def create_feeding(body: s.FeedingSessionCreate, d: Session = Depends(get_db)):
    data = body.model_dump()
    baby_id = data.pop("baby_id") or DEFAULT_BABY_ID
    logs = data.pop("logs") or []
    sess = FeedingSession(baby_id=baby_id, **data)
    d.add(sess)
    d.flush()
    for l in logs:
        d.add(FeedingLog(session_id=sess.id, feed_type=l["feed_type"], amount_ml=l.get("amount_ml")))
    d.commit()
    d.refresh(sess)
    return _feeding_payload(sess, d)


@api.put("/feedings/{item_id}", response_model=s.FeedingSessionOut, tags=["Feeding"])
def update_feeding(item_id: int, body: s.FeedingSessionCreate, d: Session = Depends(get_db)):
    sess = d.get(FeedingSession, item_id)
    if not sess:
        raise HTTPException(404, "记录不存在")
    data = body.model_dump()
    data.pop("baby_id", None)
    logs = data.pop("logs") or []
    for k, val in data.items():
        setattr(sess, k, val)
    # 奶量明细整体替换（以提交内容为准）
    for l in d.scalars(select(FeedingLog).where(FeedingLog.session_id == sess.id)).all():
        d.delete(l)
    for l in logs:
        d.add(FeedingLog(session_id=sess.id, feed_type=l["feed_type"], amount_ml=l.get("amount_ml")))
    d.commit()
    d.refresh(sess)
    return _feeding_payload(sess, d)


@api.delete("/feedings/{item_id}", tags=["Feeding"])
def delete_feeding(item_id: int, d: Session = Depends(get_db)):
    sess = d.get(FeedingSession, item_id)
    if not sess:
        raise HTTPException(404, "记录不存在")
    for l in d.scalars(select(FeedingLog).where(FeedingLog.session_id == sess.id)).all():
        d.delete(l)
    d.delete(sess)
    d.commit()
    return {"ok": True}


# ---------------- 照片 / 通用上传 ----------------
def _safe_dir_name(name: str) -> str:
    """宝宝名 → 安全目录名（去掉路径非法字符与空白）。"""
    name = re.sub(r'[\\/:*?"<>|\r\n\t ]+', "", (name or "").strip())
    return name[:20] or "宝宝"


def _meta_path(rel: str) -> str:
    """媒体相对路径 → 同目录同名 .json 说明文件路径。"""
    base = os.path.splitext(rel)[0]
    return os.path.join(UPLOAD_DIR, base + ".json")


def _write_meta(rel: str, **updates):
    """写入/合并媒体说明 JSON（本地浏览即可知类别与关联记录，不必查库）。"""
    mp = _meta_path(rel)
    meta = {}
    try:
        with open(mp, "r", encoding="utf-8") as f:
            meta = json.load(f)
    except (OSError, ValueError):
        pass
    meta.update({k: v for k, v in updates.items() if v is not None})
    try:
        with open(mp, "w", encoding="utf-8") as f:
            json.dump(meta, f, ensure_ascii=False, indent=2)
    except OSError:
        pass


def _remove_media(rel: str):
    """删除媒体文件及其 .json 说明文件（忽略不存在）。"""
    for p in (os.path.join(UPLOAD_DIR, rel), _meta_path(rel)):
        try:
            os.remove(p)
        except OSError:
            pass


def _save_media(file: UploadFile, baby_name: str, kind: str) -> str:
    """媒体归档保存：static/uploads/<宝宝名>/<YYYY-MM-DD>/<event|diaper>/<uuid>.<ext>，
    并写同名 .json 说明（kind/original_name/uploaded_at）。返回相对路径（正斜杠）。"""
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in MEDIA_EXTS:
        raise HTTPException(400, "仅支持图片/视频格式（jpg png gif webp heic mp4 mov m4v webm avi 3gp）")
    name_dir = _safe_dir_name(baby_name)
    day_dir = date.today().isoformat()
    kind_dir = kind if kind in ("event", "diaper") else "other"
    rel_dir = f"{name_dir}/{day_dir}/{kind_dir}"
    os.makedirs(os.path.join(UPLOAD_DIR, rel_dir), exist_ok=True)
    rel = f"{rel_dir}/{uuid.uuid4().hex}{ext}"
    with open(os.path.join(UPLOAD_DIR, rel), "wb") as f:
        f.write(file.file.read())
    _write_meta(
        rel,
        kind=kind_dir,
        original_name=file.filename or "",
        uploaded_at=datetime.now().isoformat(timespec="seconds"),
    )
    return rel


def migrate_uploads():
    """一次性迁移：把旧版直接放在 <宝宝名>/<日期>/ 下的媒体移入 event/diaper 子目录并补 .json。
    依据数据库归属判定类别（Photo=事件，Diaper.photo_filename=尿布）；无库记录的散文件原地保留。"""
    moved = {"event": 0, "diaper": 0}
    with db.SessionLocal() as d:
        pairs = [(p, "event") for p in d.scalars(select(Photo)).all()]
        pairs += [(x, "diaper") for x in d.scalars(select(Diaper)).all() if x.photo_filename]
        for obj, kind in pairs:
            rel = obj.filename if kind == "event" else obj.photo_filename
            # 按路径段判断是否已在类别子目录（兼容 event/x.jpg 与 x/y/event/x.jpg 两种形态）
            if not rel or kind in rel.split("/"):
                continue
            dir_part, fname = os.path.split(rel)
            # 拼接时过滤空段，避免 dir_part 为空产生前导 "/"（Windows 下会被当成盘符根路径）
            kind_rel = "/".join(p for p in (dir_part, kind, fname) if p)
            src = os.path.join(UPLOAD_DIR, rel)
            dest = os.path.join(UPLOAD_DIR, kind_rel)
            if os.path.isfile(dest):
                pass  # 上次迁移中断（文件已移走但 DB 未更新）→ 仅修 DB 指向
            elif os.path.isfile(src):
                os.makedirs(os.path.dirname(dest), exist_ok=True)
                try:
                    shutil.move(src, dest)
                except OSError as e:
                    print(f"[migrate_uploads] 移动失败 {rel}: {e!r}")
                    continue
                old_json = os.path.join(UPLOAD_DIR, dir_part, os.path.splitext(fname)[0] + ".json") if dir_part else ""
                if old_json and os.path.isfile(old_json):
                    try:
                        shutil.move(old_json, _meta_path(kind_rel))
                    except OSError:
                        pass
            else:
                continue  # 文件本就不存在（历史遗留死链）
            ts = obj.uploaded_at if kind == "event" else obj.created_at
            _write_meta(
                kind_rel,
                kind=kind,
                record_id=obj.id,
                uploaded_at=ts.isoformat(timespec="seconds") if ts else None,
            )
            if kind == "event":
                obj.filename = kind_rel
            else:
                obj.photo_filename = kind_rel
            moved[kind] += 1
        d.commit()
    if any(moved.values()):
        print(f"[migrate_uploads] 已归档迁移：事件 {moved['event']} 个，尿布 {moved['diaper']} 个")


# ---------------- 缩略图（图片按需生成 320px JPEG 缓存；img 标签无法带 token，挂 app 根路径） ----------------
THUMB_SIZE = (320, 320)
THUMB_DIR = os.path.join(UPLOAD_DIR, ".thumbs")


@app.get("/media-thumb")
def media_thumb(f: str):
    rel = (f or "").replace("\\", "/").strip("/")
    parts = [p for p in rel.split("/") if p]
    if not parts or ".." in parts:
        raise HTTPException(404)
    if os.path.splitext(parts[-1])[1].lower() not in IMAGE_EXTS:
        raise HTTPException(404)
    src = os.path.join(UPLOAD_DIR, *parts)
    if not os.path.isfile(src):
        raise HTTPException(404)
    os.makedirs(THUMB_DIR, exist_ok=True)
    tp = os.path.join(THUMB_DIR, *parts[:-1], os.path.splitext(parts[-1])[0] + ".jpg")
    if not os.path.isfile(tp):
        try:
            from PIL import Image
            os.makedirs(os.path.dirname(tp), exist_ok=True)
            img = Image.open(src)
            if img.mode not in ("RGB", "L"):
                img = img.convert("RGB")
            img.thumbnail(THUMB_SIZE)
            img.save(tp, "JPEG", quality=80)
        except Exception:
            raise HTTPException(404)   # 生成失败（如 HEIC 无解码器）：前端回退原图
    return FileResponse(tp, media_type="image/jpeg", headers={"Cache-Control": "public, max-age=86400"})


@api.post("/media/scan")
def media_scan(d: Session = Depends(get_db)):
    """扫描 uploads 目录并同步：
    1) 补建缺失的 .json 说明（已有跳过；kind 取父目录名）；
    2) 清理孤儿 .json（对应媒体已被手动删除）；
    3) 照片库同步：文件被移动→按文件名重挂；文件被手动删除→清掉死链记录；event/ 下手动拷入的新文件→入库。
    返回各类操作计数。"""
    import datetime as _dt

    stats = {"json_created": 0, "json_removed": 0, "photos_added": 0,
             "photos_removed": 0, "photos_remapped": 0, "diapers_fixed": 0}

    # 收集媒体文件：rel -> kind（父目录为 event/diaper 则取之，否则 None）
    files = {}
    for root, dirs, names in os.walk(UPLOAD_DIR):
        dirs[:] = [x for x in dirs if not x.startswith(".")]
        for n in names:
            if os.path.splitext(n)[1].lower() not in MEDIA_EXTS:
                continue
            rel = os.path.relpath(os.path.join(root, n), UPLOAD_DIR).replace("\\", "/")
            segs = rel.split("/")
            files[rel] = segs[-2] if len(segs) >= 2 and segs[-2] in ("event", "diaper") else None

    # 1) 孤儿 JSON 清理
    for root, dirs, names in os.walk(UPLOAD_DIR):
        dirs[:] = [x for x in dirs if not x.startswith(".")]
        for n in names:
            if not n.endswith(".json"):
                continue
            base = os.path.join(root, n[: -len(".json")])
            if not any(os.path.isfile(base + e) for e in MEDIA_EXTS):
                try:
                    os.remove(os.path.join(root, n))
                    stats["json_removed"] += 1
                except OSError:
                    pass

    # 2) 补建缺失 JSON（已有跳过）
    for rel, kind in files.items():
        if os.path.isfile(_meta_path(rel)):
            continue
        _write_meta(
            rel, kind=kind or "unknown",
            uploaded_at=_dt.datetime.fromtimestamp(
                os.path.getmtime(os.path.join(UPLOAD_DIR, rel))).isoformat(timespec="seconds"),
        )
        stats["json_created"] += 1

    # 3) 照片库同步（event/未分类文件均可挂为事件媒体）
    ev_by_base = {}
    di_by_base = {}
    for rel, kind in files.items():
        b = os.path.basename(rel).lower()
        if kind == "diaper":
            di_by_base.setdefault(b, []).append(rel)
        else:
            ev_by_base.setdefault(b, []).append(rel)

    used = set()
    removed_ids = []
    for p in d.scalars(select(Photo)).all():
        f = p.filename or ""
        if os.path.isfile(os.path.join(UPLOAD_DIR, f)):
            used.add(f)
            continue
        cands = [r for r in ev_by_base.get(os.path.basename(f).lower(), []) if r != f and r not in used]
        if cands:
            p.filename = cands[0]
            used.add(cands[0])
            _write_meta(cands[0], kind="event", record_id=p.id)
            stats["photos_remapped"] += 1
        else:
            removed_ids.append(p.id)
            d.delete(p)
            stats["photos_removed"] += 1
    # 新入库：event/ 下未被库引用的文件
    for rel, kind in files.items():
        if kind != "event" or rel in used:
            continue
        mt = _dt.datetime.fromtimestamp(os.path.getmtime(os.path.join(UPLOAD_DIR, rel)))
        obj = Photo(baby_id=DEFAULT_BABY_ID, filename=rel, caption="", uploaded_at=mt)
        d.add(obj)
        d.flush()
        used.add(rel)
        _write_meta(rel, kind="event", record_id=obj.id)
        stats["photos_added"] += 1
    # 清理事件对已删照片的引用
    if removed_ids:
        for m in d.scalars(select(Milestone).where(Milestone.photo_id.in_(removed_ids))).all():
            m.photo_id = None

    # 4) 尿布照片引用同步（按文件名重挂或清空；尿布记录本身不动）
    for x in d.scalars(select(Diaper)).all():
        f = x.photo_filename or ""
        if not f:
            continue
        if os.path.isfile(os.path.join(UPLOAD_DIR, f)):
            continue
        cands = [r for r in di_by_base.get(os.path.basename(f).lower(), []) if r != f]
        if cands:
            x.photo_filename = cands[0]
            _write_meta(cands[0], kind="diaper", record_id=x.id)
        else:
            x.photo_filename = ""
        stats["diapers_fixed"] += 1

    d.commit()
    return stats


# ---------------- Excel 导入（追加写入，不覆盖现有数据） ----------------
@api.post("/import/excel")
async def import_excel(file: UploadFile = File(...), d: Session = Depends(get_db)):
    """导入此前「导出 Excel」生成的文件：按自然键去重后追加写入（忽略 ID 列、不覆盖），
    用于换机迁移或向空库恢复，也可安全地重复导入。照片/视频文件不随 Excel 走，事件照片ID 置空。"""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(await file.read()), data_only=True)
    except Exception:
        raise HTTPException(400, "无法解析该文件，请使用本系统导出的 Excel")

    def pv(v):
        if isinstance(v, str):
            v = v.strip()
            return v if v not in ("", "—", "None") else None
        return v

    def col(r, i):
        """安全取第 i 列并清洗；行不够长时返回 None（openpyxl 行长度随实际数据浮动）。"""
        return pv(r[i]) if len(r) > i else None

    def pdt(v):
        v = pv(v)
        if v is None:
            return None
        if isinstance(v, datetime):
            return v
        if isinstance(v, date):
            return datetime(v.year, v.month, v.day)
        s = str(v).replace("T", " ")
        for fmt_ in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M"):
            try:
                return datetime.strptime(s[:19], fmt_)
            except ValueError:
                pass
        return None

    def pdate(v):
        dt = pdt(v)
        if dt:
            return dt.date()
        v = pv(v)
        if isinstance(v, date):
            return v
        if v:
            try:
                return date.fromisoformat(str(v)[:10])
            except ValueError:
                pass
        return None

    def pnum(v):
        v = pv(v)
        if v is None:
            return None
        try:
            return float(v)
        except (TypeError, ValueError):
            return None

    def pbool(v):
        return v in (True, 1, "1", "是", "true", "True", "TRUE")

    def rows(sheet):
        # 事件表兼容新旧名称（新版导出为「事件」，旧版导出文件仍叫「里程碑」）
        if sheet == "事件" and "事件" not in wb.sheetnames:
            sheet = "里程碑"
        if sheet not in wb.sheetnames:
            return []
        return [r for r in wb[sheet].iter_rows(min_row=2, values_only=True)
                if any(pv(c) is not None for c in r)]

    counts = {}
    skipped = {}
    bid = DEFAULT_BABY_ID
    label_type = {v: k for k, v in FEED_TYPE_LABEL.items()}

    def exist_keys(model, attrs):
        """库中已有记录的自然键集合（用于导入去重）。"""
        return {tuple(getattr(rec, a) for a in attrs)
                for rec in d.scalars(select(model)).all()}

    # 喂养：自然键 = 开始时间
    n = sk = 0
    seen = exist_keys(FeedingSession, ["start_time"])
    for r in rows("喂养"):   # [ID, 开始, 结束, 奶量明细, 总量ml, 备注]
        start = pdt(col(r, 1))
        if not start:
            continue
        if (start,) in seen:
            sk += 1
            continue
        seen.add((start,))
        sess = FeedingSession(baby_id=bid, start_time=start,
                              end_time=pdt(col(r, 2)), note=col(r, 5) or "")
        d.add(sess)
        d.flush()
        desc = col(r, 3) or ""
        for part in desc.split("+"):
            part = part.strip()
            if not part:
                continue
            m = re.match(r"^(.*?)(?:\s*([\d.]+)\s*ml)?$", part)
            label = (m.group(1) or "").strip()
            if label in label_type:
                ft = label_type[label]
            elif "母乳" in label:
                ft = "breast"
            elif "奶粉" in label:
                ft = "formula"
            else:
                ft = "mixed"
            d.add(FeedingLog(session_id=sess.id, feed_type=ft,
                             amount_ml=float(m.group(2)) if m.group(2) else None))
        n += 1
    counts["喂养"] = n
    if sk:
        skipped["喂养"] = sk

    # 睡眠：自然键 = 开始时间
    n = sk = 0
    seen = exist_keys(Sleep, ["start_time"])
    for r in rows("睡眠"):   # [ID, 开始, 结束, 备注]
        start = pdt(col(r, 1))
        if not start:
            continue
        if (start,) in seen:
            sk += 1
            continue
        seen.add((start,))
        d.add(Sleep(baby_id=bid, start_time=start,
                    end_time=pdt(col(r, 2)), note=col(r, 3) or ""))
        n += 1
    counts["睡眠"] = n
    if sk:
        skipped["睡眠"] = sk

    # 尿布：自然键 = (时间, 类型)
    n = sk = 0
    seen = exist_keys(Diaper, ["time", "diaper_type"])
    for r in rows("尿布"):   # [ID, 类型, 便便描述, 照片, 备注, 时间]
        if col(r, 1) is None and col(r, 2) is None:
            continue
        t = pdt(col(r, 5))
        dtype = col(r, 1) or "wet"
        if (t, dtype) in seen:
            sk += 1
            continue
        seen.add((t, dtype))
        d.add(Diaper(baby_id=bid, diaper_type=dtype,
                     stool_desc=col(r, 2) or "", photo_filename=col(r, 3) or "",
                     note=col(r, 4) or "", time=t))
        n += 1
    counts["尿布"] = n
    if sk:
        skipped["尿布"] = sk

    # 体温：自然键 = (时间, 数值)
    n = sk = 0
    seen = exist_keys(Temperature, ["time", "value"])
    for r in rows("体温"):   # [ID, 体温℃, 时间, 备注]
        val = pnum(col(r, 1))
        t = pdt(col(r, 2))
        if val is None or not t:
            continue
        if (t, val) in seen:
            sk += 1
            continue
        seen.add((t, val))
        d.add(Temperature(baby_id=bid, value=val, time=t, note=col(r, 3) or ""))
        n += 1
    counts["体温"] = n
    if sk:
        skipped["体温"] = sk

    # 用药：自然键 = (时间, 药名)
    n = sk = 0
    seen = exist_keys(Medication, ["time", "med_name"])
    for r in rows("用药"):   # [ID, 药名, 剂量, 服药时间, 下次服药, 备注]
        name = col(r, 1)
        t = pdt(col(r, 3))
        if not name or not t:
            continue
        if (t, name) in seen:
            sk += 1
            continue
        seen.add((t, name))
        d.add(Medication(baby_id=bid, med_name=name, dose=col(r, 2) or "",
                         time=t, next_dose_at=pdt(col(r, 4)), note=col(r, 5) or ""))
        n += 1
    counts["用药"] = n
    if sk:
        skipped["用药"] = sk

    # 疫苗：自然键 = (名称, 接种日)
    n = sk = 0
    seen = exist_keys(Vaccine, ["name", "date"])
    for r in rows("疫苗"):   # [ID, 名称, 接种日, 应种日, 已种, 备注]
        name = col(r, 1)
        if not name:
            continue
        vd = pdate(col(r, 2))
        if (name, vd) in seen:
            sk += 1
            continue
        seen.add((name, vd))
        d.add(Vaccine(baby_id=bid, name=name, date=vd,
                      due_date=pdate(col(r, 3)), done=pbool(col(r, 4)),
                      note=col(r, 5) or ""))
        n += 1
    counts["疫苗"] = n
    if sk:
        skipped["疫苗"] = sk

    # 生长：自然键 = (日期, 体重, 身高, 头围)
    n = sk = 0
    seen = exist_keys(Growth, ["record_date", "weight_kg", "height_cm", "head_cm"])
    for r in rows("生长"):   # [ID, 日期, 体重kg, 身高cm, 头围cm, 备注]
        rd = pdate(col(r, 1))
        if not rd:
            continue
        wk, hc, cc = pnum(col(r, 2)), pnum(col(r, 3)), pnum(col(r, 4))
        if (rd, wk, hc, cc) in seen:
            sk += 1
            continue
        seen.add((rd, wk, hc, cc))
        d.add(Growth(baby_id=bid, record_date=rd, weight_kg=wk, height_cm=hc, head_cm=cc,
                     note=col(r, 5) or ""))
        n += 1
    counts["生长"] = n
    if sk:
        skipped["生长"] = sk

    # 事件：自然键 = (日期, 内容)；兼容旧表名「里程碑」
    n = sk = 0
    seen = exist_keys(Milestone, ["date", "content"])
    for r in rows("事件"):   # [ID, 日期, 内容, 照片ID]
        rd = pdate(col(r, 1))
        content = col(r, 2)
        if not rd or not content:
            continue
        if (rd, content) in seen:
            sk += 1
            continue
        seen.add((rd, content))
        d.add(Milestone(baby_id=bid, date=rd, content=content, photo_id=None))
        n += 1
    counts["事件"] = n
    if sk:
        skipped["事件"] = sk

    d.commit()
    return {"ok": True, "counts": counts, "skipped": skipped}


def _register_orphan_photos(d: Session):
    """收编本地手工放入的媒体：扫描 <宝宝名>/ 下不在库里的图片/视频，
    移入所在目录的 event/ 子文件夹并注册进照片墙（文件名保持原样，不改名）。
    diaper 目录不收编（无法对应到具体尿布记录）。照片时间取文件修改时间。"""
    baby = d.get(Baby, DEFAULT_BABY_ID)
    name_dir = _safe_dir_name(baby.name if baby else "")
    root = os.path.join(UPLOAD_DIR, name_dir)
    if not os.path.isdir(root):
        return
    known = set(d.scalars(select(Photo.filename)).all())
    for dirpath, _dirnames, filenames in os.walk(root):
        rel_dir = os.path.relpath(dirpath, UPLOAD_DIR).replace(os.sep, "/")
        segs = rel_dir.split("/")
        if "diaper" in segs:
            continue
        for fname in filenames:
            ext = os.path.splitext(fname)[1].lower()
            if ext not in MEDIA_EXTS:
                continue
            rel = f"{rel_dir}/{fname}"
            if rel in known:
                continue
            # 不在 event/ 内的（如直接放在日期目录根部）先移入同级 event/
            if segs[-1] != "event":
                new_dir = os.path.join(UPLOAD_DIR, rel_dir, "event")
                os.makedirs(new_dir, exist_ok=True)
                try:
                    shutil.move(os.path.join(dirpath, fname), os.path.join(new_dir, fname))
                except OSError:
                    continue
                rel = f"{rel_dir}/event/{fname}"
            if rel in known:
                continue
            mtime = os.path.getmtime(os.path.join(UPLOAD_DIR, rel))
            obj = Photo(
                baby_id=DEFAULT_BABY_ID,
                filename=rel,
                caption="",
                uploaded_at=datetime.fromtimestamp(mtime, tz=timezone.utc).replace(tzinfo=None),
            )
            d.add(obj)
            d.commit()
            d.refresh(obj)
            known.add(rel)
            _write_meta(
                rel,
                kind="event",
                record_id=obj.id,
                original_name=fname,
                uploaded_at=datetime.fromtimestamp(mtime).isoformat(timespec="seconds"),
            )


@api.post("/photos", response_model=s.PhotoOut)
def upload_photo(
    baby_id: int = Form(None),
    caption: str = Form(""),
    file: UploadFile = File(...),
    d: Session = Depends(get_db),
):
    baby = d.get(Baby, baby_id or DEFAULT_BABY_ID)
    rel = _save_media(file, baby.name if baby else "", "event")
    obj = Photo(baby_id=DEFAULT_BABY_ID, filename=rel, caption=caption)
    d.add(obj)
    d.commit()
    d.refresh(obj)
    _write_meta(rel, kind="event", record_id=obj.id)
    return obj


@api.get("/photos", response_model=list[s.PhotoOut])
def list_photos(d: Session = Depends(get_db)):
    _register_orphan_photos(d)
    return d.scalars(select(Photo).order_by(Photo.uploaded_at.desc())).all()


@api.put("/photos/{item_id}", response_model=s.PhotoOut)
def update_photo(item_id: int, body: dict, d: Session = Depends(get_db)):
    """编辑照片/视频备注。"""
    obj = d.get(Photo, item_id)
    if not obj:
        raise HTTPException(404, "记录不存在")
    if isinstance(body.get("caption"), str):
        obj.caption = body["caption"]
    d.commit()
    d.refresh(obj)
    return obj


@api.delete("/photos/{item_id}", tags=["Photo"])
def delete_photo(item_id: int, d: Session = Depends(get_db)):
    obj = d.get(Photo, item_id)
    if not obj:
        raise HTTPException(404, "记录不存在")
    # 引用了该照片的里程碑改为无配图
    for m in d.scalars(select(Milestone).where(Milestone.photo_id == item_id)).all():
        m.photo_id = None
    _remove_media(obj.filename)
    d.delete(obj)
    d.commit()
    return {"ok": True}


@api.post("/upload")
def upload_file(file: UploadFile = File(...), d: Session = Depends(get_db)):
    """通用媒体上传：按 宝宝名/日期/diaper 归档，仅返回保存后的相对路径（尿布照片等用，不建数据库记录）。"""
    baby = d.get(Baby, DEFAULT_BABY_ID)
    rel = _save_media(file, baby.name if baby else "", "diaper")
    return {"filename": rel}


# ---------------- 概览汇总 ----------------
@api.get("/summary", response_model=s.SummaryOut)
def summary(baby_id: int = None, d: Session = Depends(get_db)):
    baby_id = baby_id or DEFAULT_BABY_ID
    today = date.today()
    start = datetime(today.year, today.month, today.day)
    end = start + timedelta(days=1)

    feed = d.execute(
        select(
            func.count(FeedingSession.id),
            func.coalesce(
                func.sum(
                    func.cast(
                        func.strftime("%s", FeedingSession.end_time) - func.strftime("%s", FeedingSession.start_time),
                        type_=db.Float,
                    ) / 60.0
                ),
                0,
            ),
        ).where(and_(FeedingSession.baby_id == baby_id, FeedingSession.start_time >= start, FeedingSession.start_time < end))
    ).first()
    feed_ml = d.execute(
        select(func.coalesce(func.sum(FeedingLog.amount_ml), 0))
        .select_from(FeedingLog)
        .join(FeedingSession, FeedingLog.session_id == FeedingSession.id)
        .where(and_(FeedingSession.baby_id == baby_id, FeedingSession.start_time >= start, FeedingSession.start_time < end))
    ).scalar()

    slp = d.execute(
        select(
            func.count(Sleep.id),
            func.coalesce(
                func.sum(
                    func.cast(
                        func.strftime("%s", Sleep.end_time) - func.strftime("%s", Sleep.start_time),
                        type_=db.Float,
                    ) / 60.0
                ),
                0,
            ),
        ).where(and_(Sleep.baby_id == baby_id, Sleep.start_time >= start, Sleep.start_time < end))
    ).first()

    dip = d.execute(
        select(func.count(Diaper.id), func.coalesce(func.sum(func.cast(Diaper.diaper_type == "poop", db.Integer)), 0))
        .where(and_(Diaper.baby_id == baby_id, Diaper.created_at >= start, Diaper.created_at < end))
    ).first()

    return s.SummaryOut(
        baby_id=baby_id,
        date=today,
        feeding_count=feed[0] or 0,
        feeding_total_ml=float(feed_ml or 0),
        feeding_total_min=int(feed[1] or 0),
        sleep_count=slp[0] or 0,
        sleep_total_min=int(slp[1] or 0),
        diaper_count=dip[0] or 0,
        diaper_poop=int(dip[1] or 0),
    )


def _norm_vax(name: str) -> str:
    """疫苗名称规范化：去空格/括号/「疫苗」字样，便于计划项与用户记录宽松匹配。
    例：「乙肝 疫苗第1针」与「乙肝第1针」匹配；数字保留，第1/第2不会混。"""
    import re
    return re.sub(r"[\s（）()·\-—_疫苗]", "", str(name or ""))


# ---------------- 提醒（疫苗 + 服药） ----------------
@api.get("/reminders")
def reminders(baby_id: int = None, d: Session = Depends(get_db)):
    baby_id = baby_id or DEFAULT_BABY_ID
    baby = d.get(Baby, baby_id)
    today = date.today()
    vaccines = []
    if baby:
        done_recs = d.scalars(select(Vaccine).where(Vaccine.baby_id == baby_id, Vaccine.done.is_(True))).all()
        done_map = {_norm_vax(v.name): v.id for v in done_recs}
        for item in VACCINE_SCHEDULE:
            due = baby.birthday + timedelta(days=int(round(item["month"] * 30.4375)))
            days_left = (due - today).days
            rec_id = done_map.get(_norm_vax(item["name"]))
            if rec_id:
                status = "已完成"
            elif days_left < 0:
                status = "已到/逾期"
            elif days_left <= 30:
                status = "即将到期"
            else:
                status = "未到"
            vaccines.append({
                "name": item["name"], "due_date": due.isoformat(),
                "days_left": days_left, "status": status,
                "record_id": rec_id,
            })
        # 用户自定义、且不在计划表中的待接种项（可就地删除）
        manual = d.scalars(select(Vaccine).where(Vaccine.baby_id == baby_id)).all()
        sched_norms = {_norm_vax(i["name"]) for i in VACCINE_SCHEDULE}
        for v in manual:
            if _norm_vax(v.name) not in sched_norms and not v.done:
                due = v.due_date or v.date
                days_left = (due - today).days if due else None
                vaccines.append({
                    "name": v.name + "（自定义）",
                    "due_date": due.isoformat() if due else None,
                    "days_left": days_left, "status": "自定义待办",
                    "record_id": v.id,
                })
        vaccines.sort(key=lambda x: (x["due_date"] is None, x["due_date"] or ""))

    meds = []
    now = datetime.now()
    for m in d.scalars(select(Medication).where(Medication.next_dose_at.isnot(None)).order_by(Medication.next_dose_at)).all():
        nd = m.next_dose_at
        if nd < now:
            status = "待服(逾期)"
        elif (nd - now).total_seconds() <= 86400:
            status = "24小时内"
        else:
            status = "计划中"
        meds.append({
            "id": m.id, "med_name": m.med_name, "dose": m.dose,
            "next_dose_at": nd.isoformat(), "status": status,
        })

    # 未来日期的「事件」（里程碑）自动转为提醒：date 晚于今天的都进入提醒列表
    events = []
    if baby:
        for m in d.scalars(select(Milestone).where(Milestone.baby_id == baby_id)).all():
            if m.date and m.date > today:
                events.append({
                    "name": (m.content or "").strip().splitlines()[0][:60] if m.content else "事件",
                    "date": m.date.isoformat(),
                    "days_left": (m.date - today).days,
                    "status": "事件提醒",
                })
        events.sort(key=lambda x: x["date"])
    return {"vaccines": vaccines, "medications": meds, "events": events}


# ---------------- 数据导出 ----------------
@api.get("/export/excel")
def export_excel(d: Session = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    baby = d.get(Baby, DEFAULT_BABY_ID)

    def sheet(name, headers, rows):
        ws = wb.create_sheet(title=name)
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        return ws

    # 宝宝信息
    ws0 = wb.active
    ws0.title = "宝宝信息"
    ws0.append(["名称", "出生日期", "性别", "备注", "出生天数"])
    if baby:
        ws0.append([baby.name, str(baby.birthday), baby.gender, baby.note, (date.today() - baby.birthday).days])

    def rows_of(model, cols):
        items = d.scalars(select(model).order_by(model.id)).all()
        return [[getattr(x, c) for c in cols] for x in items]

    def feeding_rows():
        out = []
        for x in d.scalars(select(FeedingSession).order_by(FeedingSession.id)).all():
            logs = d.scalars(select(FeedingLog).where(FeedingLog.session_id == x.id).order_by(FeedingLog.id)).all()
            desc = " + ".join(
                FEED_TYPE_LABEL.get(l.feed_type, l.feed_type) + (f" {l.amount_ml}ml" if l.amount_ml is not None else "")
                for l in logs
            ) or "—"
            total = sum(l.amount_ml or 0 for l in logs)
            out.append([x.id, x.start_time, x.end_time, desc, round(total, 1), x.note])
        return out

    sheet("喂养", ["ID", "开始", "结束", "奶量明细", "总量ml", "备注"], feeding_rows())
    sheet("睡眠", ["ID", "开始", "结束", "备注"],
          rows_of(Sleep, ["id", "start_time", "end_time", "note"]))
    sheet("尿布", ["ID", "类型", "便便描述", "照片", "备注", "时间"],
          [[x.id, x.diaper_type, x.stool_desc, x.photo_filename, x.note, x.created_at] for x in d.scalars(select(Diaper).order_by(Diaper.id)).all()])
    sheet("体温", ["ID", "体温℃", "时间", "备注"],
          rows_of(Temperature, ["id", "value", "time", "note"]))
    sheet("用药", ["ID", "药名", "剂量", "服药时间", "下次服药", "备注"],
          [[x.id, x.med_name, x.dose, x.time, x.next_dose_at, x.note] for x in d.scalars(select(Medication).order_by(Medication.id)).all()])
    sheet("疫苗", ["ID", "名称", "接种日", "应种日", "已种", "备注"],
          rows_of(Vaccine, ["id", "name", "date", "due_date", "done", "note"]))
    sheet("生长", ["ID", "日期", "体重kg", "身高cm", "头围cm", "备注"],
          rows_of(Growth, ["id", "record_date", "weight_kg", "height_cm", "head_cm", "note"]))
    sheet("事件", ["ID", "日期", "内容", "照片ID"],
          rows_of(Milestone, ["id", "date", "content", "photo_id"]))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=baoji_export.xlsx"},
    )


@api.get("/export/pdf")
def export_pdf(d: Session = Depends(get_db)):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    CN = "STSong-Light"
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title="宝记 数据导出")
    styles = getSampleStyleSheet()
    h = ParagraphStyle("h", parent=styles["Title"], fontName=CN, fontSize=18)
    body = ParagraphStyle("b", parent=styles["Normal"], fontName=CN, fontSize=9)
    elems = []

    baby = d.get(Baby, DEFAULT_BABY_ID)
    elems.append(Paragraph("宝记 · 宝宝成长数据", h))
    if baby:
        elems.append(Paragraph(
            f"姓名：{baby.name}　性别：{baby.gender}　出生：{baby.birthday}　出生第 {(date.today() - baby.birthday).days} 天", body))
    elems.append(Spacer(1, 6))

    def table(title, headers, rows):
        elems.append(Paragraph(title, body))
        data = [[Paragraph(str(c), body) for c in headers]]
        for r in rows:
            data.append([Paragraph("" if v is None else str(v), body) for v in r])
        t = Table(data, repeatRows=1, colWidths=[None] * len(headers))
        t.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), CN),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8f1fb")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        elems.append(t)
        elems.append(Spacer(1, 8))

    def feeding_table_rows():
        out = []
        for x in d.scalars(select(FeedingSession).order_by(FeedingSession.id)).all():
            logs = d.scalars(select(FeedingLog).where(FeedingLog.session_id == x.id).order_by(FeedingLog.id)).all()
            desc = " + ".join(
                FEED_TYPE_LABEL.get(l.feed_type, l.feed_type) + (f" {l.amount_ml}ml" if l.amount_ml is not None else "")
                for l in logs
            ) or "—"
            out.append([x.start_time, x.end_time, desc, sum(l.amount_ml or 0 for l in logs), x.note])
        return out

    table("喂养记录", ["开始", "结束", "奶量明细", "总量ml", "备注"], feeding_table_rows())
    table("睡眠记录", ["开始", "结束", "备注"],
          [[x.start_time, x.end_time, x.note] for x in d.scalars(select(Sleep).order_by(Sleep.id)).all()])
    table("尿布记录", ["类型", "便便描述", "照片", "备注", "时间"],
          [[x.diaper_type, x.stool_desc, x.photo_filename, x.note, x.created_at] for x in d.scalars(select(Diaper).order_by(Diaper.id)).all()])
    table("体温记录", ["体温℃", "时间", "备注"],
          [[x.value, x.time, x.note] for x in d.scalars(select(Temperature).order_by(Temperature.id)).all()])
    table("用药记录", ["药名", "剂量", "服药时间", "下次", "备注"],
          [[x.med_name, x.dose, x.time, x.next_dose_at, x.note] for x in d.scalars(select(Medication).order_by(Medication.id)).all()])
    table("疫苗记录", ["名称", "接种日", "应种日", "已种", "备注"],
          [[x.name, x.date, x.due_date, x.done, x.note] for x in d.scalars(select(Vaccine).order_by(Vaccine.id)).all()])
    table("生长记录", ["日期", "体重kg", "身高cm", "头围cm", "备注"],
          [[x.record_date, x.weight_kg, x.height_cm, x.head_cm, x.note] for x in d.scalars(select(Growth).order_by(Growth.id)).all()])
    table("事件记录", ["日期", "内容"],
          [[x.date, x.content] for x in d.scalars(select(Milestone).order_by(Milestone.id)).all()])

    doc.build(elems)
    buf.seek(0)
    return Response(
        content=buf.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=baoji_export.pdf"},
    )


# ---------------- 挂载路由 & 静态资源 ----------------
app.include_router(api, prefix="/api")


@app.get("/")
def index():
    return FileResponse(os.path.join(BASE_DIR, "static", "index.html"))


app.mount("/static", StaticFiles(directory=os.path.join(BASE_DIR, "static")), name="static")


@app.on_event("startup")
def _startup():
    global DEFAULT_BABY_ID
    db.init_db()
    ensure_default_baby()
    DEFAULT_BABY_ID = get_single_baby_id()
    migrate_uploads()


if __name__ == "__main__":
    import argparse
    import uvicorn

    parser = argparse.ArgumentParser(description="宝记 局域网宝宝数据收集")
    parser.add_argument("--host", default=os.environ.get("BABYWEB_HOST", "0.0.0.0"))
    parser.add_argument("--port", type=int, default=int(os.environ.get("BABYWEB_PORT", "8000")),
                        help="端口；设为 0 由系统自动分配空闲端口")
    parser.add_argument("--no-auth", action="store_true", help="关闭访问口令验证（仅限可信局域网）")
    args = parser.parse_args()

    NO_AUTH = args.no_auth or NO_AUTH
    if args.port == 0:
        sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        sock.bind(("0.0.0.0", 0))
        port = sock.getsockname()[1]
        sock.close()
    else:
        port = args.port

    print(f"【宝记】已启动：http://{args.host}:{port}  （本机请用 http://localhost:{port}）")
    if NO_AUTH:
        print("【宝记】注意：已关闭访问验证，请确保仅在可信局域网内使用。")
    uvicorn.run(app, host=args.host, port=port)
